"""Verify and hash the generated Phase E annotation deliverables."""

from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

EXPECTED_ITEMS = 360
OUTPUT_SUBDIR = "outputs/phase_e_v2_180jobs"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    package = root / "human_validation"
    output = package / OUTPUT_SUBDIR
    workbooks = [
        output / "Human_A_annotations.xlsx",
        output / "Human_B_annotations.xlsx",
    ]
    artifacts = workbooks + [
        output / "Familiarization_set.xlsx",
        output / "ANNOTATOR_INSTRUCTIONS.pdf",
    ]
    for path in artifacts:
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"missing or empty output: {path}")

    from openpyxl import load_workbook

    for path in workbooks:
        with zipfile.ZipFile(path) as archive:
            if "xl/workbook.xml" not in set(archive.namelist()):
                raise RuntimeError(f"invalid annotation workbook: {path}")
        workbook = load_workbook(path, read_only=True)
        if workbook.sheetnames != ["Instructions", "Annotator", "Annotation"]:
            raise RuntimeError(f"{path}: unexpected sheets {workbook.sheetnames}")
        sheet = workbook["Annotation"]
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        expected_header = [
            "Item", "Prompt", "Model response",
            "Your score (0-3)", "Borderline", "Cannot read",
        ]
        if header != expected_header:
            raise RuntimeError(f"{path}: unexpected header {header}")
        rows = list(sheet.iter_rows(min_row=2, max_col=6, values_only=True))
        if len(rows) != EXPECTED_ITEMS:
            raise RuntimeError(f"{path}: expected {EXPECTED_ITEMS} items, found {len(rows)}")
        if [row[0] for row in rows] != list(range(1, EXPECTED_ITEMS + 1)):
            raise RuntimeError(f"{path}: item numbers are not 1..{EXPECTED_ITEMS}")
        if any(row[3] is not None for row in rows):
            raise RuntimeError(f"{path}: a score cell is pre-filled")
        if any(not row[1] or not row[2] for row in rows):
            raise RuntimeError(f"{path}: an item is missing prompt or response text")

    if not (output / "ANNOTATOR_INSTRUCTIONS.pdf").read_bytes().startswith(b"%PDF"):
        raise RuntimeError("instruction document is not a PDF")

    manifest_path = package / "HUMAN_VALIDATION_MANIFEST.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest["response_items_per_annotator"] != EXPECTED_ITEMS:
        raise RuntimeError("manifest and workbooks disagree on item count")
    manifest["output_hashes"] = {
        str(path.relative_to(root)).replace("\\", "/"): sha256_file(path) for path in artifacts
    }
    manifest["protocol_hashes"] = {
        str(path.relative_to(root)).replace("\\", "/"): sha256_file(path)
        for path in (
            package / "ANNOTATOR_INSTRUCTIONS.md",
            package / "ADJUDICATION_PROTOCOL.md",
            package / "analyze_returned_labels.py",
        )
    }
    manifest_path.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    print(
        f"Human-validation package PASS: {manifest['selected_pair_model_jobs']} jobs, "
        f"{EXPECTED_ITEMS} items per annotator, {len(artifacts)} deliverables hashed"
    )


if __name__ == "__main__":
    main()
